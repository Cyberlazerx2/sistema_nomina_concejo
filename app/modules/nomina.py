import datetime
from decimal import ROUND_HALF_UP, Decimal
import io
import json
import re
from app.database import get_db_connection
from app.modules.payroll_engine import calcular_quincena_trabajador
from config import Config
from flask import (
    Blueprint,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

nomina_bp = Blueprint('nomina', __name__)


# -------------------------------------------------------------------------
# 1. HISTORIAL Y TABLERO DE NÓMINAS
# -------------------------------------------------------------------------
@nomina_bp.route('/historial')
def historial():
  if 'usuario_id' not in session:
    return redirect(url_for('auth.login'))

  roles_permitidos = [
      'analista_rrhh',
      'directora_rrhh',
      'tecnico_informatica',
      'presidente',
  ]
  if session.get('rol') not in roles_permitidos:
    flash('Acceso denegado al historial de nóminas.', 'danger')
    return redirect(url_for('auth.dashboard'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      sql = """
                SELECT p.*, u.nombre_usuario AS analista_nombre, tn.nombre AS tipo_nomina_nombre
                FROM periodo_nomina p
                LEFT JOIN usuario u ON p.elaborada_por = u.usuario_id
                LEFT JOIN tipo_nomina tn ON p.tipo_nomina = tn.codigo
                ORDER BY p.periodo_id DESC;
            """
      cursor.execute(sql)
      periodos = cursor.fetchall()

      cursor.execute(
          'SELECT * FROM tipo_nomina WHERE activo = true ORDER BY nombre ASC;'
      )
      tipos_raw = cursor.fetchall()
      tipos_nomina = [(r['codigo'], r['nombre']) for r in tipos_raw]

      tipos_info = {}
      for r in tipos_raw:
        es_manual = bool(
            r.get('requiere_monto_manual', 0)
            or 'bono' in r['codigo'].lower()
            or r['codigo'] == 'cestaticket'
            or r.get('periodicidad') == 'especial'
        )
        tipos_info[r['codigo']] = {
            'nombre': r['nombre'],
            'requiere_monto_manual': es_manual,
            'aplica_leyes_sociales': bool(r['aplica_leyes_sociales']),
            'periodicidad': r['periodicidad'],
        }

      return render_template(
          'nomina/nomina_historial.html',
          periodos=periodos,
          tipos_nomina=tipos_nomina,
          tipos_info_json=json.dumps(tipos_info),
      )
  finally:
    conexion.close()


# -------------------------------------------------------------------------
# 2. CATÁLOGO DE TIPOS DE NÓMINA
# -------------------------------------------------------------------------
@nomina_bp.route('/tipos')
def tipos_catalogo():
  if 'usuario_id' not in session or session.get('rol') not in [
      'directora_rrhh',
      'tecnico_informatica',
  ]:
    flash(
        'Acceso restringido: Solo la Dirección de RRHH puede parametrizar tipos'
        ' de nómina.',
        'danger',
    )
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      sql = """
                SELECT tn.*,
                       COUNT(DISTINCT t.trabajador_id) AS total_adscritos,
                       COUNT(DISTINCT pn.periodo_id) AS total_periodos
                FROM tipo_nomina tn
                LEFT JOIN trabajador t ON tn.codigo = t.tipo_personal AND t.activo = true
                LEFT JOIN periodo_nomina pn ON tn.codigo = pn.tipo_nomina
                GROUP BY tn.codigo, tn.nombre, tn.periodicidad, tn.aplica_leyes_sociales, tn.aplica_a_todos, tn.activo
                ORDER BY tn.nombre ASC;
            """
      cursor.execute(sql)
      tipos = cursor.fetchall()
      return render_template('nomina/nomina_tipos.html', tipos=tipos)
  finally:
    conexion.close()


@nomina_bp.route('/tipos/guardar', methods=['POST'])
def guardar_tipo():
  if 'usuario_id' not in session or session.get('rol') not in [
      'directora_rrhh',
      'tecnico_informatica',
  ]:
    flash('Acceso no autorizado.', 'danger')
    return redirect(url_for('nomina.historial'))

  codigo_original = request.form.get('codigo_original', '').strip()
  nombre = request.form.get('nombre', '').strip()
  periodicidad = request.form.get('periodicidad', 'quincenal')
  aplica_leyes = 1 if request.form.get('aplica_leyes_sociales') == '1' else 0
  aplica_todos = 1 if request.form.get('aplica_a_todos') == '1' else 0
  requiere_manual = (
      1 if request.form.get('requiere_monto_manual') == '1' else 0
  )

  if not nombre:
    flash('El nombre del tipo de nómina es obligatorio.', 'warning')
    return redirect(url_for('nomina.tipos_catalogo'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      if codigo_original:
        cursor.execute(
            """
                    UPDATE tipo_nomina 
                    SET nombre = %s, periodicidad = %s, aplica_leyes_sociales = %s, 
                        aplica_a_todos = %s, requiere_monto_manual = %s
                    WHERE codigo = %s;
                """,
            (
                nombre,
                periodicidad,
                aplica_leyes,
                aplica_todos,
                requiere_manual,
                codigo_original,
            ),
        )
        flash(f"Tipo de nómina '{nombre}' actualizado.", 'success')
      else:
        codigo = re.sub(
            r'[^a-zA-Z0-9_]',
            '',
            request.form.get('codigo', '').strip().lower().replace(' ', '_'),
        )
        if not codigo:
          codigo = re.sub(
              r'[^a-zA-Z0-9_]', '', nombre.lower().replace(' ', '_')
          )[:25]

        cursor.execute(
            """
                    INSERT INTO tipo_nomina (codigo, nombre, periodicidad, aplica_leyes_sociales, aplica_a_todos, requiere_monto_manual, activo)
                    VALUES (%s, %s, %s, %s, %s, %s, 1);
                """,
            (
                codigo,
                nombre,
                periodicidad,
                aplica_leyes,
                aplica_todos,
                requiere_manual,
            ),
        )

        if requiere_manual:
          cursor.execute(
              'SELECT concepto_id FROM concepto WHERE codigo = %s;',
              (codigo.upper()[:20],),
          )
          if not cursor.fetchone():
            cursor.execute(
                """
                            INSERT INTO concepto (codigo, nombre, tipo, metodo, valor, base_calculo, obligatorio, fundamento, activo)
                            VALUES (%s, %s, 'asignacion', 'fijo', 0.0000, 'sueldo_base', 0, 'Asignación por Decreto Municipal', 1);
                        """,
                (codigo.upper()[:20], nombre[:80]),
            )

        flash(f"Nuevo tipo de nómina '{nombre}' configurado.", 'success')

      conexion.commit()
  except Exception as e:
    conexion.rollback()
    flash(f'Error al registrar tipo de nómina: {e}', 'danger')
  finally:
    conexion.close()

  return redirect(url_for('nomina.tipos_catalogo'))


@nomina_bp.route('/tipos/<codigo>/toggle', methods=['POST'])
def toggle_tipo(codigo):
  if 'usuario_id' not in session or session.get('rol') not in [
      'directora_rrhh',
      'tecnico_informatica',
  ]:
    flash('Acceso no autorizado.', 'danger')
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      cursor.execute(
          'UPDATE tipo_nomina SET activo = NOT activo WHERE codigo = %s;',
          (codigo,),
      )
      conexion.commit()
      flash(f"Estatus del tipo de nómina '{codigo}' actualizado.", 'info')
  finally:
    conexion.close()

  return redirect(url_for('nomina.tipos_catalogo'))


# -------------------------------------------------------------------------
# 3. GENERAR PRE-NÓMINA
# -------------------------------------------------------------------------
@nomina_bp.route('/generar', methods=['GET', 'POST'])
def generar():
  roles_permitidos = ['analista_rrhh', 'directora_rrhh']
  if 'usuario_id' not in session or session.get('rol') not in roles_permitidos:
    flash('Acceso no autorizado para elaborar nóminas.', 'danger')
    return redirect(url_for('auth.dashboard'))

  if request.method != 'POST':
    return redirect(url_for('nomina.historial'))

  tipo_personal = request.form.get('tipo_personal')
  fecha_inicio = request.form.get('fecha_inicio')
  fecha_fin = request.form.get('fecha_fin')
  descripcion = request.form.get('descripcion', '').strip()
  monto_bono_raw = request.form.get('monto_bono', '').strip()

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      cursor.execute(
          'SELECT * FROM tipo_nomina WHERE codigo = %s;', (tipo_personal,)
      )
      regla_tipo = cursor.fetchone()

      if not regla_tipo:
        flash(
            'El tipo de nómina seleccionado no está configurado en el sistema.',
            'danger',
        )
        return redirect(url_for('nomina.historial'))

      aplica_a_todos = bool(regla_tipo['aplica_a_todos'])
      aplica_leyes_sociales = bool(regla_tipo['aplica_leyes_sociales'])
      periodicidad = regla_tipo['periodicidad']

      es_monto_manual = bool(
          regla_tipo.get('requiere_monto_manual', 0)
          or 'bono' in tipo_personal.lower()
          or regla_tipo['periodicidad'] == 'especial'
      )

      monto_manual_val = None
      if es_monto_manual:
        try:
          monto_manual_val = Decimal(monto_bono_raw.replace(',', '.'))
          if monto_manual_val <= Decimal('0.00'):
            flash(
                'Debe indicar un monto válido mayor a cero para esta nómina'
                ' extraordinaria.',
                'warning',
            )
            return redirect(url_for('nomina.historial'))
        except Exception:
          flash(
              'El monto manual ingresado no posee un formato numérico válido.',
              'danger',
          )
          return redirect(url_for('nomina.historial'))

      if not descripcion:
        descripcion = (
            f"Nómina {regla_tipo['nombre']} ({fecha_inicio} al {fecha_fin})"
        )

      # Resolver concepto para asignación manual y evitar Error 1452
      concepto_manual_obj = None
      if es_monto_manual:
        cod_buscar = (
            'CESTATICKET'
            if tipo_personal == 'cestaticket'
            else tipo_personal.upper()[:20]
        )
        cursor.execute(
            'SELECT * FROM concepto WHERE codigo = %s;', (cod_buscar,)
        )
        concepto_manual_obj = cursor.fetchone()

        if not concepto_manual_obj:
          cursor.execute(
              "SELECT * FROM concepto WHERE codigo = 'BONO_EXTRA';"
          )
          concepto_manual_obj = cursor.fetchone()

        if not concepto_manual_obj:
          cursor.execute(
              """
                        INSERT INTO concepto (codigo, nombre, tipo, metodo, valor, base_calculo, obligatorio, fundamento, activo)
                        VALUES (%s, %s, 'asignacion', 'fijo', 0.0000, 'sueldo_base', 0, 'Asignación por Decreto Municipal', 1);
                    """,
              (cod_buscar, regla_tipo['nombre'][:80]),
          )
          nuevo_id = cursor.lastrowid
          concepto_manual_obj = {
              'concepto_id': nuevo_id,
              'codigo': cod_buscar,
              'nombre': regla_tipo['nombre'][:80],
              'tipo': 'asignacion',
          }

      if aplica_a_todos:
        cursor.execute("""
                    SELECT trabajador_id, cedula, nombres, apellidos, sueldo_base, 
                           hijos, nivel_instruccion, tipo_personal 
                    FROM trabajador WHERE activo = true;
                """)
      else:
        cursor.execute(
            """
                    SELECT trabajador_id, cedula, nombres, apellidos, sueldo_base, 
                           hijos, nivel_instruccion, tipo_personal 
                    FROM trabajador WHERE tipo_personal = %s AND activo = true;
                """,
            (tipo_personal,),
        )
      plantilla = cursor.fetchall()

      if not plantilla:
        flash(
            'No hay funcionarios activos para la nómina:'
            f" {regla_tipo['nombre']}",
            'warning',
        )
        return redirect(url_for('nomina.historial'))

      cursor.execute('SELECT * FROM concepto WHERE activo = true;')
      conceptos = cursor.fetchall()

      resultados = []
      gran_total_asig = Decimal('0.00')
      gran_total_ded = Decimal('0.00')
      gran_total_neto = Decimal('0.00')

      sql_ind = """
                SELECT c.*, ct.valor AS valor_individual FROM concepto_trabajador ct 
                JOIN concepto c ON ct.concepto_id = c.concepto_id 
                WHERE ct.trabajador_id = %s AND ct.activo = true 
                AND (ct.fecha_fin IS NULL OR ct.fecha_fin >= CURDATE());
            """

      for trab in plantilla:
        cursor.execute(sql_ind, (trab['trabajador_id'],))
        conceptos_individuales = cursor.fetchall()
        for ci in conceptos_individuales:
          ci['valor'] = ci['valor_individual']

        try:
          calc = calcular_quincena_trabajador(
              trabajador=trab,
              conceptos_catalogo=conceptos,
              conceptos_individuales=conceptos_individuales,
              solo_cestaticket=(tipo_personal == 'cestaticket'),
              aplica_leyes_sociales=aplica_leyes_sociales,
              es_monto_manual=es_monto_manual,
              monto_manual=monto_manual_val,
              concepto_asignacion_manual=concepto_manual_obj,
          )
        except TypeError:
          calc = calcular_quincena_trabajador(
              trabajador=trab,
              conceptos_catalogo=conceptos,
              conceptos_individuales=conceptos_individuales,
              solo_cestaticket=(tipo_personal == 'cestaticket'),
              aplica_leyes_sociales=aplica_leyes_sociales,
              es_monto_manual=es_monto_manual,
              monto_manual=monto_manual_val,
          )

        resultados.append(calc)
        gran_total_asig += calc['total_asignaciones']
        gran_total_ded += calc['total_deducciones']
        gran_total_neto += calc['neto_pagar']

      sql_periodo = """
                INSERT INTO periodo_nomina (
                    descripcion, tipo_nomina, periodicidad, fecha_inicio, fecha_fin, 
                    estado, elaborada_por, fecha_elaboracion, 
                    total_asignaciones, total_deducciones, total_neto
                ) VALUES (%s, %s, %s, %s, %s, 'borrador', %s, %s, %s, %s, %s);
            """
      ahora = datetime.datetime.now()
      cursor.execute(
          sql_periodo,
          (
              descripcion[:80],
              tipo_personal,
              periodicidad,
              fecha_inicio,
              fecha_fin,
              session['usuario_id'],
              ahora,
              gran_total_asig,
              gran_total_ded,
              gran_total_neto,
          ),
      )
      periodo_id = cursor.lastrowid

      for r in resultados:
        cursor.execute(
            """
                    INSERT INTO nomina_trabajador (
                        periodo_id, trabajador_id, dias_trabajados, 
                        total_asignaciones, total_deducciones, neto_pagar
                    ) VALUES (%s, %s, %s, %s, %s, %s);
                """,
            (
                periodo_id,
                r['trabajador_id'],
                r['dias_trabajados'],
                r['total_asignaciones'],
                r['total_deducciones'],
                r['neto_pagar'],
            ),
        )
        nomina_trabajador_id = cursor.lastrowid

        if r['detalles']:
          detalles_datos = [
              (
                  nomina_trabajador_id,
                  d['concepto_id'],
                  d['codigo_concepto'],
                  d['descripcion'][:80],
                  d['tipo'],
                  d['monto'],
              )
              for d in r['detalles']
          ]
          cursor.executemany(
              """
                        INSERT INTO nomina_detalle (
                            nomina_trabajador_id, concepto_id, codigo_concepto, descripcion, tipo, monto
                        ) VALUES (%s, %s, %s, %s, %s, %s);
                    """,
              detalles_datos,
          )

      conexion.commit()
      flash(
          '¡Pre-nómina generada exitosamente en BORRADOR! Periodo'
          f' #{periodo_id}',
          'success',
      )
      return redirect(url_for('nomina.detalle', periodo_id=periodo_id))

  except Exception as e:
    conexion.rollback()
    flash(f'Error al procesar la nómina: {e}', 'danger')
  finally:
    conexion.close()

  return redirect(url_for('nomina.historial'))


# -------------------------------------------------------------------------
# 4. DETALLE DEL PERIODO DE NÓMINA
# -------------------------------------------------------------------------
@nomina_bp.route('/detalle/<int:periodo_id>')
def detalle(periodo_id):
  if 'usuario_id' not in session:
    return redirect(url_for('auth.login'))

  roles_permitidos = [
      'analista_rrhh',
      'directora_rrhh',
      'tecnico_informatica',
      'presidente',
  ]
  if session.get('rol') not in roles_permitidos:
    flash('Acceso denegado.', 'danger')
    return redirect(url_for('auth.dashboard'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      sql_p = """
                SELECT p.*, 
                       u1.nombre_usuario AS analista_nom,
                       u2.nombre_usuario AS directora_nom,
                       u3.nombre_usuario AS presidente_nom,
                       tn.nombre AS tipo_nomina_nombre
                FROM periodo_nomina p
                LEFT JOIN usuario u1 ON p.elaborada_por = u1.usuario_id
                LEFT JOIN usuario u2 ON p.revisada_por = u2.usuario_id
                LEFT JOIN usuario u3 ON p.aprobada_por = u3.usuario_id
                LEFT JOIN tipo_nomina tn ON p.tipo_nomina = tn.codigo
                WHERE p.periodo_id = %s;
            """
      cursor.execute(sql_p, (periodo_id,))
      periodo = cursor.fetchone()

      if not periodo:
        flash('El periodo solicitado no existe.', 'warning')
        return redirect(url_for('nomina.historial'))

      sql_trab = """
                SELECT nt.*, t.cedula, t.nombres, t.apellidos, t.banco, t.numero_cuenta, 
                       t.fecha_ingreso, t.sueldo_base, c.nombre AS cargo_nombre,
                       d.nombre AS departamento_nombre
                FROM nomina_trabajador nt
                INNER JOIN trabajador t ON nt.trabajador_id = t.trabajador_id
                INNER JOIN cargo c ON t.cargo_id = c.cargo_id
                LEFT JOIN departamento d ON c.departamento_id = d.departamento_id
                WHERE nt.periodo_id = %s
                ORDER BY t.apellidos ASC, t.nombres ASC;
            """
      cursor.execute(sql_trab, (periodo_id,))
      trabajadores = cursor.fetchall()

      for trab in trabajadores:
        sql_det = """
                    SELECT codigo_concepto, descripcion, tipo, monto 
                    FROM nomina_detalle 
                    WHERE nomina_trabajador_id = %s
                    ORDER BY tipo ASC, monto DESC;
                """
        cursor.execute(sql_det, (trab['nomina_trabajador_id'],))
        trab['desglose'] = cursor.fetchall()

      return render_template(
          'nomina/nomina_detalle.html',
          periodo=periodo,
          trabajadores=trabajadores,
      )
  finally:
    conexion.close()


# -------------------------------------------------------------------------
# 5. REVISAR NÓMINA (Directora)
# -------------------------------------------------------------------------
@nomina_bp.route('/revisar/<int:periodo_id>', methods=['POST'])
def revisar(periodo_id):
  if 'usuario_id' not in session or session.get('rol') != 'directora_rrhh':
    flash(
        'Acceso restringido: Solo la Directora de RRHH puede validar nóminas.',
        'danger',
    )
    return redirect(url_for('nomina.historial'))

  accion = request.form.get('accion')
  motivo = request.form.get('observaciones_rechazo', '').strip()

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      ahora = datetime.datetime.now()
      if accion == 'validar':
        sql = """
                    UPDATE periodo_nomina 
                    SET estado = 'revisada', revisada_por = %s, fecha_revision = %s,
                        observaciones_rechazo = NULL 
                    WHERE periodo_id = %s;
                """
        cursor.execute(sql, (session['usuario_id'], ahora, periodo_id))
        conexion.commit()
        flash(
            f'Periodo #{periodo_id} VALIDADO. Listo para firma de Presidencia.',
            'success',
        )

      elif accion == 'rechazar':
        sql = """
                    UPDATE periodo_nomina 
                    SET estado = 'rechazada', revisada_por = NULL, fecha_revision = NULL,
                        observaciones_rechazo = %s 
                    WHERE periodo_id = %s;
                """
        cursor.execute(
            sql, (motivo or 'Rechazada por observaciones técnicas.', periodo_id)
        )
        conexion.commit()
        flash(f'Periodo #{periodo_id} marcado como RECHAZADO.', 'warning')
  finally:
    conexion.close()

  return redirect(url_for('nomina.detalle', periodo_id=periodo_id))


# -------------------------------------------------------------------------
# 6. APROBAR NÓMINA (Directora o Presidente)
# -------------------------------------------------------------------------
@nomina_bp.route('/aprobar/<int:periodo_id>', methods=['POST'])
def aprobar(periodo_id):
  roles_autorizados = ['directora_rrhh', 'presidente']
  if (
      'usuario_id' not in session
      or session.get('rol') not in roles_autorizados
  ):
    flash('Acceso restringido para registrar la aprobación oficial.', 'danger')
    return redirect(url_for('nomina.historial'))

  accion = request.form.get('accion')
  motivo = request.form.get('observaciones_rechazo', '').strip()

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      ahora = datetime.datetime.now()
      if accion == 'aprobar':
        sql = """
                    UPDATE periodo_nomina 
                    SET estado = 'aprobada', aprobada_por = %s, fecha_aprobacion = %s,
                        observaciones_rechazo = NULL 
                    WHERE periodo_id = %s AND estado IN ('borrador', 'revisada');
                """
        cursor.execute(sql, (session['usuario_id'], ahora, periodo_id))
        conexion.commit()
        flash(f'Periodo #{periodo_id} APROBADO oficialmente.', 'success')

      elif accion == 'devolver':
        sql = """
                    UPDATE periodo_nomina 
                    SET estado = 'rechazada', revisada_por = NULL, fecha_revision = NULL,
                        observaciones_rechazo = %s 
                    WHERE periodo_id = %s;
                """
        cursor.execute(
            sql,
            (
                motivo or 'Devuelta para correcciones administrativas.',
                periodo_id,
            ),
        )
        conexion.commit()
        flash(f'Periodo #{periodo_id} devuelto a estado rechazado.', 'warning')
  finally:
    conexion.close()

  return redirect(url_for('nomina.detalle', periodo_id=periodo_id))


# -------------------------------------------------------------------------
# 7. RETROCEDER ESTADO DE NÓMINA (Apertura de Período para Correcciones)
# -------------------------------------------------------------------------
@nomina_bp.route('/retroceder/<int:periodo_id>', methods=['POST'])
def retroceder(periodo_id):
  roles_permitidos = [
      'directora_rrhh',
      'analista_rrhh',
      'tecnico_informatica',
  ]
  if 'usuario_id' not in session or session.get('rol') not in roles_permitidos:
    flash('Acceso restringido para modificar el ciclo de nómina.', 'danger')
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      sql = """
                UPDATE periodo_nomina 
                SET estado = 'borrador', 
                    revisada_por = NULL, fecha_revision = NULL,
                    aprobada_por = NULL, fecha_aprobacion = NULL,
                    observaciones_rechazo = NULL
                WHERE periodo_id = %s;
            """
      cursor.execute(sql, (periodo_id,))
      conexion.commit()
      flash(
          f'Período #{periodo_id} reabierto en estado BORRADOR. Listo para'
          ' recálculo o ajustes.',
          'info',
      )
  finally:
    conexion.close()

  return redirect(url_for('nomina.detalle', periodo_id=periodo_id))


# -------------------------------------------------------------------------
# 8. RECALCULAR PRE-NÓMINA (Actualiza según salarios y plantilla actuales)
# -------------------------------------------------------------------------
@nomina_bp.route('/recalcular/<int:periodo_id>', methods=['POST'])
def recalcular(periodo_id):
  roles_permitidos = ['directora_rrhh', 'analista_rrhh']
  if 'usuario_id' not in session or session.get('rol') not in roles_permitidos:
    flash('Acceso no autorizado para recalcular nóminas.', 'danger')
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      cursor.execute(
          'SELECT * FROM periodo_nomina WHERE periodo_id = %s;', (periodo_id,)
      )
      periodo = cursor.fetchone()

      if not periodo:
        flash('El período solicitado no existe.', 'warning')
        return redirect(url_for('nomina.historial'))

      cursor.execute(
          'SELECT * FROM tipo_nomina WHERE codigo = %s;',
          (periodo['tipo_nomina'],),
      )
      regla_tipo = cursor.fetchone()

      aplica_a_todos = bool(regla_tipo['aplica_a_todos']) if regla_tipo else False
      aplica_leyes_sociales = (
          bool(regla_tipo['aplica_leyes_sociales']) if regla_tipo else True
      )
      es_monto_manual = (
          bool(
              regla_tipo.get('requiere_monto_manual', 0)
              or 'bono' in periodo['tipo_nomina'].lower()
              or periodo['periodicidad'] == 'especial'
          )
          if regla_tipo
          else False
      )

      monto_manual_val = None
      if es_monto_manual:
        cursor.execute(
            """
                    SELECT monto FROM nomina_detalle 
                    WHERE nomina_trabajador_id IN (SELECT nomina_trabajador_id FROM nomina_trabajador WHERE periodo_id = %s)
                      AND tipo = 'asignacion' LIMIT 1;
                """,
            (periodo_id,),
        )
        res_m = cursor.fetchone()
        monto_manual_val = (
            Decimal(str(res_m['monto'])) if res_m else Decimal('0.00')
        )

      concepto_manual_obj = None
      if es_monto_manual:
        cod_buscar = (
            'CESTATICKET'
            if periodo['tipo_nomina'] == 'cestaticket'
            else periodo['tipo_nomina'].upper()[:20]
        )
        cursor.execute(
            'SELECT * FROM concepto WHERE codigo = %s;', (cod_buscar,)
        )
        concepto_manual_obj = cursor.fetchone()
        if not concepto_manual_obj:
          cursor.execute(
              "SELECT * FROM concepto WHERE codigo = 'BONO_EXTRA';"
          )
          concepto_manual_obj = cursor.fetchone()

      if aplica_a_todos:
        cursor.execute("""
                    SELECT trabajador_id, cedula, nombres, apellidos, sueldo_base, 
                           hijos, nivel_instruccion, tipo_personal 
                    FROM trabajador WHERE activo = true;
                """)
      else:
        cursor.execute(
            """
                    SELECT trabajador_id, cedula, nombres, apellidos, sueldo_base, 
                           hijos, nivel_instruccion, tipo_personal 
                    FROM trabajador WHERE tipo_personal = %s AND activo = true;
                """,
            (periodo['tipo_nomina'],),
        )
      plantilla = cursor.fetchall()

      cursor.execute('SELECT * FROM concepto WHERE activo = true;')
      conceptos = cursor.fetchall()

      resultados = []
      gran_total_asig = Decimal('0.00')
      gran_total_ded = Decimal('0.00')
      gran_total_neto = Decimal('0.00')

      sql_ind = """
                SELECT c.*, ct.valor AS valor_individual FROM concepto_trabajador ct 
                JOIN concepto c ON ct.concepto_id = c.concepto_id 
                WHERE ct.trabajador_id = %s AND ct.activo = true 
                AND (ct.fecha_fin IS NULL OR ct.fecha_fin >= CURDATE());
            """

      for trab in plantilla:
        cursor.execute(sql_ind, (trab['trabajador_id'],))
        conceptos_individuales = cursor.fetchall()
        for ci in conceptos_individuales:
          ci['valor'] = ci['valor_individual']

        try:
          calc = calcular_quincena_trabajador(
              trabajador=trab,
              conceptos_catalogo=conceptos,
              conceptos_individuales=conceptos_individuales,
              solo_cestaticket=(periodo['tipo_nomina'] == 'cestaticket'),
              aplica_leyes_sociales=aplica_leyes_sociales,
              es_monto_manual=es_monto_manual,
              monto_manual=monto_manual_val,
              concepto_asignacion_manual=concepto_manual_obj,
          )
        except TypeError:
          calc = calcular_quincena_trabajador(
              trabajador=trab,
              conceptos_catalogo=conceptos,
              conceptos_individuales=conceptos_individuales,
              solo_cestaticket=(periodo['tipo_nomina'] == 'cestaticket'),
              aplica_leyes_sociales=aplica_leyes_sociales,
              es_monto_manual=es_monto_manual,
              monto_manual=monto_manual_val,
          )

        resultados.append(calc)
        gran_total_asig += calc['total_asignaciones']
        gran_total_ded += calc['total_deducciones']
        gran_total_neto += calc['neto_pagar']

      # Limpiar registros anteriores del período
      cursor.execute(
          'SELECT nomina_trabajador_id FROM nomina_trabajador WHERE periodo_id'
          ' = %s;',
          (periodo_id,),
      )
      ids_nt = [r['nomina_trabajador_id'] for r in cursor.fetchall()]
      if ids_nt:
        formato_in = ','.join(['%s'] * len(ids_nt))
        cursor.execute(
            f'DELETE FROM nomina_detalle WHERE nomina_trabajador_id IN'
            f' ({formato_in});',
            ids_nt,
        )
        cursor.execute(
            'DELETE FROM nomina_trabajador WHERE periodo_id = %s;', (periodo_id,)
        )

      for r in resultados:
        cursor.execute(
            """
                    INSERT INTO nomina_trabajador (
                        periodo_id, trabajador_id, dias_trabajados, 
                        total_asignaciones, total_deducciones, neto_pagar
                    ) VALUES (%s, %s, %s, %s, %s, %s);
                """,
            (
                periodo_id,
                r['trabajador_id'],
                r['dias_trabajados'],
                r['total_asignaciones'],
                r['total_deducciones'],
                r['neto_pagar'],
            ),
        )
        nomina_trabajador_id = cursor.lastrowid

        if r['detalles']:
          detalles_datos = [
              (
                  nomina_trabajador_id,
                  d['concepto_id'],
                  d['codigo_concepto'],
                  d['descripcion'][:80],
                  d['tipo'],
                  d['monto'],
              )
              for d in r['detalles']
          ]
          cursor.executemany(
              """
                        INSERT INTO nomina_detalle (
                            nomina_trabajador_id, concepto_id, codigo_concepto, descripcion, tipo, monto
                        ) VALUES (%s, %s, %s, %s, %s, %s);
                    """,
              detalles_datos,
          )

      cursor.execute(
          """
                UPDATE periodo_nomina 
                SET total_asignaciones = %s, total_deducciones = %s, total_neto = %s, estado = 'borrador'
                WHERE periodo_id = %s;
            """,
          (gran_total_asig, gran_total_ded, gran_total_neto, periodo_id),
      )

      conexion.commit()
      flash(
          f'¡Período #{periodo_id} recalculado con éxito! Nuevo total neto: Bs.'
          f' {gran_total_neto:,.2f}',
          'success',
      )
  except Exception as e:
    conexion.rollback()
    flash(f'Error al recalcular nómina: {e}', 'danger')
  finally:
    conexion.close()

  return redirect(url_for('nomina.detalle', periodo_id=periodo_id))


# -------------------------------------------------------------------------
# 9. ELIMINAR NÓMINA (Directora puede purgar nóminas de prueba o viejas)
# -------------------------------------------------------------------------
@nomina_bp.route('/eliminar/<int:periodo_id>', methods=['POST'])
def eliminar(periodo_id):
  roles_permitidos = [
      'directora_rrhh',
      'analista_rrhh',
      'tecnico_informatica',
  ]
  if 'usuario_id' not in session or session.get('rol') not in roles_permitidos:
    flash('Acceso restringido.', 'danger')
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      cursor.execute(
          'SELECT * FROM periodo_nomina WHERE periodo_id = %s;', (periodo_id,)
      )
      periodo = cursor.fetchone()

      if not periodo:
        flash('El periodo solicitado no existe.', 'warning')
        return redirect(url_for('nomina.historial'))

      es_directora_o_admin = session.get('rol') in [
          'directora_rrhh',
          'tecnico_informatica',
      ]
      if not es_directora_o_admin and periodo['estado'] not in [
          'borrador',
          'rechazada',
      ]:
        flash(
            'Acción protegida: Los analistas solo pueden eliminar borradores o'
            ' nóminas rechazadas.',
            'danger',
        )
        return redirect(url_for('nomina.historial'))

      cursor.execute(
          'SELECT nomina_trabajador_id FROM nomina_trabajador WHERE periodo_id'
          ' = %s;',
          (periodo_id,),
      )
      ids_nt = [r['nomina_trabajador_id'] for r in cursor.fetchall()]

      if ids_nt:
        formato_in = ','.join(['%s'] * len(ids_nt))
        cursor.execute(
            f'DELETE FROM nomina_detalle WHERE nomina_trabajador_id IN'
            f' ({formato_in});',
            ids_nt,
        )
        cursor.execute(
            'DELETE FROM nomina_trabajador WHERE periodo_id = %s;', (periodo_id,)
        )

      cursor.execute(
          'DELETE FROM periodo_nomina WHERE periodo_id = %s;', (periodo_id,)
      )
      conexion.commit()
      flash(
          f"El período #{periodo_id} ('{periodo['descripcion']}') fue"
          ' eliminado definitivamente.',
          'info',
      )
  except Exception as e:
    conexion.rollback()
    flash(f'Error al eliminar el período: {e}', 'danger')
  finally:
    conexion.close()

  return redirect(url_for('nomina.historial'))


# -------------------------------------------------------------------------
# 10. RECIBOS INDIVIDUALES Y EN LOTE
# -------------------------------------------------------------------------
@nomina_bp.route('/recibo/<int:nomina_trabajador_id>')
def recibo_individual(nomina_trabajador_id):
  roles_permitidos = [
      'directora_rrhh',
      'analista_rrhh',
      'tecnico_informatica',
      'presidente',
  ]
  if 'usuario_id' not in session or session.get('rol') not in roles_permitidos:
    flash('Acceso denegado.', 'danger')
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      sql = """
                SELECT nt.nomina_trabajador_id, nt.total_asignaciones, nt.total_deducciones, nt.neto_pagar,
                       t.trabajador_id, t.cedula, t.nombres, t.apellidos, t.fecha_ingreso,
                       t.banco, t.numero_cuenta, t.sueldo_base,
                       c.nombre AS cargo_nombre, d.nombre AS departamento_nombre,
                       pn.periodo_id, pn.descripcion, pn.fecha_inicio, pn.fecha_fin,
                       pn.tipo_nomina, pn.periodicidad, pn.estado
                FROM nomina_trabajador nt
                INNER JOIN periodo_nomina pn ON nt.periodo_id = pn.periodo_id
                INNER JOIN trabajador t ON nt.trabajador_id = t.trabajador_id
                INNER JOIN cargo c ON t.cargo_id = c.cargo_id
                LEFT JOIN departamento d ON c.departamento_id = d.departamento_id
                WHERE nt.nomina_trabajador_id = %s;
            """
      cursor.execute(sql, (nomina_trabajador_id,))
      recibo = cursor.fetchone()

      if not recibo:
        flash('Recibo no encontrado.', 'warning')
        return redirect(url_for('nomina.historial'))

      cursor.execute(
          """
                SELECT tipo, descripcion, monto 
                FROM nomina_detalle 
                WHERE nomina_trabajador_id = %s
                ORDER BY tipo ASC, monto DESC;
            """,
          (nomina_trabajador_id,),
      )
      desglose = cursor.fetchall()

      asignaciones = [d for d in desglose if d.get('tipo') == 'asignacion']
      deducciones = [d for d in desglose if d.get('tipo') == 'deduccion']
      fecha_emision = datetime.date.today().strftime('%d/%m/%Y')

      return render_template(
          'portal/portal_recibo_detalle.html',
          r=recibo,
          asignaciones=asignaciones,
          deducciones=deducciones,
          fecha_emision=fecha_emision,
      )
  finally:
    conexion.close()


@nomina_bp.route('/recibos-lote/<int:periodo_id>')
def recibos_lote(periodo_id):
  roles_permitidos = [
      'directora_rrhh',
      'analista_rrhh',
      'tecnico_informatica',
      'presidente',
  ]
  if 'usuario_id' not in session or session.get('rol') not in roles_permitidos:
    flash('Acceso denegado.', 'danger')
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      cursor.execute(
          'SELECT * FROM periodo_nomina WHERE periodo_id = %s;', (periodo_id,)
      )
      periodo = cursor.fetchone()

      if not periodo:
        flash('Periodo no encontrado.', 'warning')
        return redirect(url_for('nomina.historial'))

      sql = """
                SELECT nt.nomina_trabajador_id, nt.total_asignaciones, nt.total_deducciones, nt.neto_pagar,
                       t.trabajador_id, t.cedula, t.nombres, t.apellidos, t.fecha_ingreso,
                       t.banco, t.numero_cuenta, t.sueldo_base,
                       c.nombre AS cargo_nombre, d.nombre AS departamento_nombre
                FROM nomina_trabajador nt
                INNER JOIN trabajador t ON nt.trabajador_id = t.trabajador_id
                INNER JOIN cargo c ON t.cargo_id = c.cargo_id
                LEFT JOIN departamento d ON c.departamento_id = d.departamento_id
                WHERE nt.periodo_id = %s
                ORDER BY t.apellidos ASC, t.nombres ASC;
            """
      cursor.execute(sql, (periodo_id,))
      trabajadores = cursor.fetchall()

      for trab in trabajadores:
        cursor.execute(
            """
                    SELECT tipo, descripcion, monto 
                    FROM nomina_detalle 
                    WHERE nomina_trabajador_id = %s
                    ORDER BY tipo ASC, monto DESC;
                """,
            (trab['nomina_trabajador_id'],),
        )
        desglose = cursor.fetchall()
        trab['asignaciones'] = [
            d for d in desglose if d.get('tipo') == 'asignacion'
        ]
        trab['deducciones'] = [
            d for d in desglose if d.get('tipo') == 'deduccion'
        ]

      fecha_emision = datetime.date.today().strftime('%d/%m/%Y')

      return render_template(
          'nomina/nomina_recibos_lote.html',
          periodo=periodo,
          trabajadores=trabajadores,
          fecha_emision=fecha_emision,
      )
  finally:
    conexion.close()


# -------------------------------------------------------------------------
# 11. SÁBANA CONSOLIDADA
# -------------------------------------------------------------------------
@nomina_bp.route('/reporte-consolidado/<int:periodo_id>')
def reporte_consolidado(periodo_id):
  if 'usuario_id' not in session:
    return redirect(url_for('auth.login'))

  roles_permitidos = [
      'analista_rrhh',
      'directora_rrhh',
      'tecnico_informatica',
      'presidente',
  ]
  if session.get('rol') not in roles_permitidos:
    flash('Acceso denegado.', 'danger')
    return redirect(url_for('auth.dashboard'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      sql_p = """
                SELECT p.*, 
                       u1.nombre_usuario AS analista_nom,
                       u2.nombre_usuario AS directora_nom,
                       u3.nombre_usuario AS presidente_nom,
                       tn.nombre AS tipo_nomina_nombre,
                       tn.aplica_leyes_sociales,
                       tn.requiere_monto_manual
                FROM periodo_nomina p
                LEFT JOIN usuario u1 ON p.elaborada_por = u1.usuario_id
                LEFT JOIN usuario u2 ON p.revisada_por = u2.usuario_id
                LEFT JOIN usuario u3 ON p.aprobada_por = u3.usuario_id
                LEFT JOIN tipo_nomina tn ON p.tipo_nomina = tn.codigo
                WHERE p.periodo_id = %s;
            """
      cursor.execute(sql_p, (periodo_id,))
      periodo = cursor.fetchone()

      if not periodo:
        flash('El periodo solicitado no existe.', 'warning')
        return redirect(url_for('nomina.historial'))

      sql_trab = """
                SELECT nt.*, 
                       t.cedula, t.nombres, t.apellidos, t.fecha_ingreso, t.sueldo_base,
                       c.nombre AS cargo_nombre
                FROM nomina_trabajador nt
                INNER JOIN trabajador t ON nt.trabajador_id = t.trabajador_id
                INNER JOIN cargo c ON t.cargo_id = c.cargo_id
                WHERE nt.periodo_id = %s
                ORDER BY t.apellidos ASC, t.nombres ASC;
            """
      cursor.execute(sql_trab, (periodo_id,))
      trabajadores_raw = cursor.fetchall()

      es_bono_o_cesta = (
          periodo['tipo_nomina']
          in ['cestaticket', 'bono_alimentacion', 'bono_especial', 'jubilados']
          or periodo.get('aplica_leyes_sociales') == 0
          or periodo.get('requiere_monto_manual') == 1
          or 'bono' in periodo['tipo_nomina'].lower()
      )

      filas = []
      totales = {
          'sueldo_mensual': Decimal('0.00'),
          'sueldo_quincenal': Decimal('0.00'),
          'ivss': Decimal('0.00'),
          'spf': Decimal('0.00'),
          'faov': Decimal('0.00'),
          'fejp': Decimal('0.00'),
          'total_deducciones': Decimal('0.00'),
          'subtotal': Decimal('0.00'),
          'prima_prof': Decimal('0.00'),
          'prima_antig': Decimal('0.00'),
          'prima_hijos': Decimal('0.00'),
          'otras_asig': Decimal('0.00'),
          'total_asignaciones': Decimal('0.00'),
          'neto_pagar': Decimal('0.00'),
      }

      for idx, t in enumerate(trabajadores_raw, start=1):
        cursor.execute(
            """
                    SELECT codigo_concepto, tipo, monto 
                    FROM nomina_detalle 
                    WHERE nomina_trabajador_id = %s;
                """,
            (t['nomina_trabajador_id'],),
        )
        detalles = cursor.fetchall()
        conceptos = {
            d['codigo_concepto']: Decimal(str(d['monto'])) for d in detalles
        }

        if es_bono_o_cesta:
          sueldo_mensual = Decimal('0.00')
          sueldo_quincenal = Decimal('0.00')
          subtotal = Decimal('0.00')
          prima_prof = Decimal('0.00')
          prima_antig = Decimal('0.00')
          prima_hijos = Decimal('0.00')
          otras_asig = Decimal(str(t['total_asignaciones']))
        else:
          sueldo_mensual = Decimal(str(t['sueldo_base']))
          sueldo_quincenal = (sueldo_mensual / Decimal('2.00')).quantize(
              Decimal('0.01'), rounding=ROUND_HALF_UP
          )
          prima_prof = (
              conceptos.get('PRIMA_PROF_TSU', Decimal('0.00'))
              + conceptos.get('PRIMA_PROF_LIC', Decimal('0.00'))
              + conceptos.get('PRIMA_PROF_ESP', Decimal('0.00'))
          )
          prima_antig = conceptos.get('PRIMA_ANTIG', Decimal('0.00'))
          prima_hijos = conceptos.get('PRIMA_HIJO', Decimal('0.00'))
          asignaciones_principales = (
              prima_prof + prima_antig + prima_hijos + sueldo_quincenal
          )
          otras_asig = (
              Decimal(str(t['total_asignaciones'])) - asignaciones_principales
          )
          if otras_asig < Decimal('0.00'):
            otras_asig = Decimal('0.00')

        ivss = conceptos.get('IVSS', Decimal('0.00'))
        spf = conceptos.get('RPE', Decimal('0.00'))
        faov = conceptos.get('FAOV', Decimal('0.00'))
        fejp = conceptos.get('FONDO_JUB', Decimal('0.00'))
        total_ded = Decimal(str(t['total_deducciones']))

        if not es_bono_o_cesta:
          subtotal = (sueldo_quincenal - total_ded).quantize(
              Decimal('0.01'), rounding=ROUND_HALF_UP
          )

        neto = Decimal(str(t['neto_pagar']))

        fila = {
            'num': idx,
            'cedula': t['cedula'],
            'nombre_completo': f"{t['nombres']} {t['apellidos']}",
            'cargo': t['cargo_nombre'],
            'fecha_ingreso': t['fecha_ingreso'],
            'sueldo_mensual': sueldo_mensual,
            'sueldo_quincenal': sueldo_quincenal,
            'ivss': ivss,
            'spf': spf,
            'faov': faov,
            'fejp': fejp,
            'total_deducciones': total_ded,
            'subtotal': subtotal,
            'prima_prof': prima_prof,
            'prima_antig': prima_antig,
            'prima_hijos': prima_hijos,
            'otras_asig': otras_asig,
            'total_asignaciones': Decimal(str(t['total_asignaciones'])),
            'neto_pagar': neto,
        }
        filas.append(fila)

        totales['sueldo_mensual'] += sueldo_mensual
        totales['sueldo_quincenal'] += sueldo_quincenal
        totales['ivss'] += ivss
        totales['spf'] += spf
        totales['faov'] += faov
        totales['fejp'] += fejp
        totales['total_deducciones'] += total_ded
        totales['subtotal'] += subtotal
        totales['prima_prof'] += prima_prof
        totales['prima_antig'] += prima_antig
        totales['prima_hijos'] += prima_hijos
        totales['otras_asig'] += otras_asig
        totales['total_asignaciones'] += Decimal(str(t['total_asignaciones']))
        totales['neto_pagar'] += neto

      titulo_nomina = (
          periodo.get('tipo_nomina_nombre') or periodo['tipo_nomina']
      ).upper()

      return render_template(
          'nomina/nomina_consolidada.html',
          periodo=periodo,
          titulo_nomina=titulo_nomina,
          filas=filas,
          totales=totales,
          es_cestaticket=es_bono_o_cesta,
      )
  finally:
    conexion.close()


# -------------------------------------------------------------------------
# 12. EXPORTACIÓN A PATRIA TXT
# -------------------------------------------------------------------------
@nomina_bp.route('/exportar-patria/<int:periodo_id>')
def exportar_patria(periodo_id):
  roles_autorizados = [
      'directora_rrhh',
      'analista_rrhh',
      'tecnico_informatica',
  ]
  if 'usuario_id' not in session or session.get('rol') not in roles_autorizados:
    flash('Acceso restringido.', 'danger')
    return redirect(url_for('nomina.historial'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      cursor.execute(
          'SELECT * FROM periodo_nomina WHERE periodo_id = %s;', (periodo_id,)
      )
      periodo = cursor.fetchone()

      if not periodo or periodo['estado'] not in [
          'aprobada',
          'procesada_patria',
          'revisada',
      ]:
        flash(
            'Solo se pueden exportar nóminas validadas o aprobadas'
            ' formalmente.',
            'warning',
        )
        return redirect(url_for('nomina.detalle', periodo_id=periodo_id))

      sql = """
                SELECT t.cedula, t.numero_cuenta, nt.neto_pagar, t.nombres, t.apellidos
                FROM nomina_trabajador nt
                INNER JOIN trabajador t ON nt.trabajador_id = t.trabajador_id
                WHERE nt.periodo_id = %s AND nt.neto_pagar > 0
                ORDER BY t.apellidos ASC, t.nombres ASC;
            """
      cursor.execute(sql, (periodo_id,))
      registros = cursor.fetchall()

      if not registros:
        flash(
            'No hay registros de pago válidos para exportar en este período.',
            'warning',
        )
        return redirect(url_for('nomina.detalle', periodo_id=periodo_id))

      cuenta_institucion = getattr(
          Config, 'CUENTA_PAGADORA_CONCEJO', '01750071650000001149'
      )
      fecha_ref = periodo['fecha_fin']
      fecha_str = (
          fecha_ref.strftime('%Y%m%d')
          if hasattr(fecha_ref, 'strftime')
          else str(fecha_ref).replace('-', '')[:8]
      )

      total_centimos = sum(
          int(round(float(r['neto_pagar']) * 100)) for r in registros
      )
      cant_registros = len(registros)

      # Cabecera estándar (49 caracteres)
      linea_cabecera = f'{cuenta_institucion[:20]:<20}{fecha_str:0>8}{total_centimos:017d}{cant_registros:04d}'

      # Detalle bancario (54 caracteres por trabajador)
      lineas = [linea_cabecera]
      for r in registros:
        cod_transaccion = '0713'
        monto_centimos = int(round(float(r['neto_pagar']) * 100))
        monto_str = f'{monto_centimos:012d}'
        cuenta_raw = re.sub(r'\D', '', str(r['numero_cuenta'] or ''))[:20]
        cuenta_20 = f'{cuenta_raw:0>20}'
        cedula_raw = re.sub(r'\D', '', str(r['cedula'] or ''))[:10]
        cedula_10 = (
            f'{int(cedula_raw):010d}' if cedula_raw else '0000000000'
        )
        relleno = '00000000'

        lineas.append(
            f'{cod_transaccion}{monto_str}{cuenta_20}{cedula_10}{relleno}'
        )

      contenido_txt = '\r\n'.join(lineas) + '\r\n'
      buffer = io.BytesIO(contenido_txt.encode('utf-8'))

      ahora = datetime.datetime.now()
      nombre_archivo = f"TXT_BICENTENARIO_{periodo['tipo_nomina'].upper()}_P{periodo_id}_{fecha_str}.txt"

      cursor.execute(
          """
                UPDATE periodo_nomina 
                SET estado = 'procesada_patria', fecha_proceso_patria = %s, archivo_txt_patria = %s 
                WHERE periodo_id = %s;
            """,
          (ahora, nombre_archivo, periodo_id),
      )
      conexion.commit()

      return send_file(
          buffer,
          as_attachment=True,
          download_name=nombre_archivo,
          mimetype='text/plain',
      )
  finally:
    conexion.close()


# -------------------------------------------------------------------------
# 13. EXPORTACIÓN A EXCEL
# -------------------------------------------------------------------------
@nomina_bp.route('/exportar-excel-bono/<int:periodo_id>')
def exportar_excel_bono(periodo_id):
  if 'usuario_id' not in session:
    return redirect(url_for('auth.login'))

  conexion = get_db_connection()
  try:
    with conexion.cursor() as cursor:
      cursor.execute(
          """
                SELECT p.*, tn.nombre AS tipo_nomina_nombre 
                FROM periodo_nomina p
                LEFT JOIN tipo_nomina tn ON p.tipo_nomina = tn.codigo
                WHERE p.periodo_id = %s;
            """,
          (periodo_id,),
      )
      periodo = cursor.fetchone()

      if not periodo:
        flash('Período no encontrado.', 'warning')
        return redirect(url_for('nomina.historial'))

      sql = """
                SELECT nt.neto_pagar, t.cedula, t.nombres, t.apellidos, t.tipo_personal
                FROM nomina_trabajador nt
                INNER JOIN trabajador t ON nt.trabajador_id = t.trabajador_id
                WHERE nt.periodo_id = %s
                ORDER BY FIELD(t.tipo_personal, 'alto_nivel', 'libre_remocion', 'empleado_fijo', 'empleado_contratado', 'obrero_fijo', 'obrero_contratado', 'jubilados'),
                         t.apellidos ASC, t.nombres ASC;
            """
      cursor.execute(sql, (periodo_id,))
      trabajadores = cursor.fetchall()

      wb = openpyxl.Workbook()
      ws = wb.active
      ws.title = 'RELACION NOMINA'

      ws.merge_cells('A1:D1')
      ws['A1'] = 'CONCEJO MUNICIPAL DE INDEPENDENCIA'
      ws['A1'].font = Font(name='Calibri', size=13, bold=True)
      ws['A1'].alignment = Alignment(horizontal='center')

      ws.merge_cells('A2:D2')
      ws['A2'] = (
          f"NOMINA {periodo.get('tipo_nomina_nombre', periodo['tipo_nomina']).upper()}"
      )
      ws['A2'].font = Font(name='Calibri', size=11, bold=True)
      ws['A2'].alignment = Alignment(horizontal='center')

      ws.merge_cells('A3:D3')
      ws['A3'] = (
          f"PERÍODO: {periodo['descripcion']} (HASTA {periodo['fecha_fin']})"
      )
      ws['A3'].font = Font(name='Calibri', size=10, italic=True)
      ws['A3'].alignment = Alignment(horizontal='center')

      headers = ['Nº', 'CÉDULA Nº', 'NOMBRES Y APELLIDOS', 'TOTAL A PAGAR']
      ws.append([])
      ws.append(headers)

      fill_hdr = PatternFill(
          start_color='D9E1F2', end_color='D9E1F2', fill_type='solid'
      )
      font_hdr = Font(name='Calibri', size=10, bold=True)
      for col_num in range(1, 5):
        cell = ws.cell(row=5, column=col_num)
        cell.fill = fill_hdr
        cell.font = font_hdr
        cell.alignment = Alignment(horizontal='center')

      etiquetas_grupos = {
          'alto_nivel': 'CONCEJALES',
          'libre_remocion': 'PERSONAL DIRECTIVO',
          'empleado_fijo': 'EMPLEADOS FIJOS',
          'empleado_contratado': 'EMPLEADOS CONTRATADOS',
          'obrero_fijo': 'OBREROS FIJOS',
          'obrero_contratado': 'OBREROS CONTRATADOS',
          'jubilados': 'JUBILADOS Y PENSIONADOS',
      }

      current_row = 6
      contador_total = 1
      grupo_actual = None

      for t in trabajadores:
        grupo = t['tipo_personal']
        if grupo != grupo_actual:
          grupo_actual = grupo
          ws.cell(
              row=current_row,
              column=1,
              value=etiquetas_grupos.get(grupo, grupo.upper()),
          ).font = Font(bold=True)
          current_row += 1

        ws.cell(row=current_row, column=1, value=contador_total).alignment = (
            Alignment(horizontal='center')
        )
        ws.cell(row=current_row, column=2, value=str(t['cedula'])).alignment = (
            Alignment(horizontal='center')
        )
        ws.cell(
            row=current_row, column=3, value=f"{t['nombres']} {t['apellidos']}"
        )
        cell_monto = ws.cell(
            row=current_row, column=4, value=float(t['neto_pagar'])
        )
        cell_monto.number_format = '#,##0.00'
        cell_monto.alignment = Alignment(horizontal='right')

        current_row += 1
        contador_total += 1

      ws.cell(row=current_row, column=3, value='TOTAL GENERAL').font = Font(
          bold=True
      )
      cell_tot = ws.cell(
          row=current_row, column=4, value=float(periodo['total_neto'] or 0)
      )
      cell_tot.font = Font(bold=True)
      cell_tot.number_format = '#,##0.00'

      ws.column_dimensions['A'].width = 8
      ws.column_dimensions['B'].width = 16
      ws.column_dimensions['C'].width = 40
      ws.column_dimensions['D'].width = 20

      buffer = io.BytesIO()
      wb.save(buffer)
      buffer.seek(0)

      nombre_excel = f'NOMINA_{periodo_id}.xlsx'
      return send_file(
          buffer,
          as_attachment=True,
          download_name=nombre_excel,
          mimetype=(
              'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
          ),
      )
  finally:
    conexion.close()